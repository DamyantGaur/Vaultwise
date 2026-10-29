"""Private user documents: contracts, company policies shared with the user, confidential files.

Text is extracted locally, clause/section headings are detected, and the result is stored as Markdown in
`upload_dir` (outside the cloud-synced project folder by default). Ids are prefixed "user-" so private
documents can be listed, filtered and deleted without touching the company policy set used for evaluation.
"""

from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from pathlib import Path

UPLOAD_PREFIX = "user-"
CATEGORIES = ("Contract", "Company policy", "Confidential document", "Other")
SUPPORTED_UPLOADS = (".pdf", ".docx", ".txt", ".md")
MAX_UPLOAD_BYTES = 10 * 1024 * 1024

# "5. Termination", "5.2 Notice Period", "Clause 7 - Confidentiality", "ARTICLE IV: Benefits",
# "SECTION 3 COMPENSATION" or a short ALL-CAPS line such as "NON-COMPETE".
_CLAUSE_RE = re.compile(
    r"^((clause|section|article)\s+[\divxlc]+[.:)\-\s]|\d{1,2}(\.\d{1,2})*[.)]?\s+[A-Z])", re.IGNORECASE)


def _pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader
        except ImportError:
            PdfReader = None
    if PdfReader is not None:
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    import fitz  # PyMuPDF
    with fitz.open(stream=data, filetype="pdf") as doc:
        return "\n".join(page.get_text() for page in doc)


def _docx_text(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    lines = []
    for p in d.paragraphs:
        text = p.text.strip()
        if text and p.style is not None and p.style.name.lower().startswith("heading"):
            text = "## " + text
        lines.append(text)
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_UPLOADS:
        raise ValueError(f"Unsupported file type '{ext}'. Use one of: {', '.join(SUPPORTED_UPLOADS)}")
    if ext == ".pdf":
        text = _pdf_text(data)
    elif ext == ".docx":
        text = _docx_text(data)
    else:
        text = data.decode("utf-8", errors="replace")
    if not text.strip():
        raise ValueError("No text could be extracted (a scanned PDF needs OCR first).")
    return text


def is_clause_heading(line: str) -> bool:
    s = line.strip()
    if not s or len(s) > 90 or s.endswith((",", ";")):
        return False
    if _CLAUSE_RE.match(s) and len(s.split()) <= 10:
        return True
    letters = [c for c in s if c.isalpha()]
    return len(letters) >= 4 and all(c.isupper() for c in letters) and len(s.split()) <= 8


def to_markdown(title: str, text: str, category: str = "Other", added: str | None = None) -> str:
    """Turn extracted text into the Markdown layout the chunker expects: clause and section headings
    become '##' sections, so retrieved chunks (and citations) point at a specific clause."""
    category = category if category in CATEGORIES else "Other"
    added = added or f"{datetime.now(timezone.utc):%Y-%m-%d}"  # fixed in evaluations, so the index is identical
    out = [f"# {title}", "", f"Type: {category} | Private | Added: {added}", ""]
    for raw in text.replace("\r\n", "\n").split("\n"):
        line = raw.strip()
        if line.startswith("#"):
            out.append("## " + line.lstrip("#").strip())
        elif is_clause_heading(line):
            out.append("## " + line.rstrip(":"))
        else:
            out.append(line)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip() + "\n"


def slugify(text: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return slug[:50].strip("-") or "document"


def save_user_document(upload_dir: str | Path, title: str, text: str, category: str = "Other",
                       added: str | None = None) -> str:
    upload_dir = Path(upload_dir)
    upload_dir.mkdir(parents=True, exist_ok=True)
    base = UPLOAD_PREFIX + slugify(title)
    doc_id, n = base, 2
    while (upload_dir / f"{doc_id}.md").exists():
        doc_id, n = f"{base}-{n}", n + 1
    (upload_dir / f"{doc_id}.md").write_text(to_markdown(title.strip() or doc_id, text, category, added),
                                             encoding="utf-8")
    return doc_id


def _user_path(upload_dir: str | Path, doc_id: str) -> Path:
    if not doc_id.startswith(UPLOAD_PREFIX) or not re.fullmatch(r"[a-z0-9-]+", doc_id):
        raise ValueError("Only your own documents (ids starting with 'user-') can be opened or deleted.")
    return Path(upload_dir) / f"{doc_id}.md"


def read_user_document(upload_dir: str | Path, doc_id: str) -> dict | None:
    """Return {doc_id, title, category, added, text} for one private document, or None."""
    path = _user_path(upload_dir, doc_id)
    if not path.exists():
        return None
    text = path.read_text(encoding="utf-8")
    title = re.search(r"^#\s+(.+)$", text, re.MULTILINE)
    meta = re.search(r"^Type:\s*([^|\n]+?)\s*\|.*?Added:\s*(\S+)", text, re.MULTILINE)
    return {"doc_id": doc_id, "title": title.group(1).strip() if title else doc_id,
            "category": meta.group(1).strip() if meta else "Other", "added": meta.group(2) if meta else "",
            "text": text}


def delete_user_document(upload_dir: str | Path, doc_id: str) -> bool:
    path = _user_path(upload_dir, doc_id)
    if not path.exists():
        return False
    path.unlink()
    return True
