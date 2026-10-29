import io
from pathlib import Path

import docx
import pytest
from fastapi.testclient import TestClient

from conftest import FakeClient
from vaultwise.api import create_app
from vaultwise.documents import (delete_user_document, extract_text, is_clause_heading, save_user_document,
                                   to_markdown)
from vaultwise.ingest import split_sections
from vaultwise.llm import Embedder
from vaultwise.pipeline import RAGPipeline

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "samples" / "sample-employment-contract.txt"


@pytest.fixture
def ucfg(cfg, tmp_path):
    return dict(cfg, upload_dir=str(tmp_path / "uploads"))


def make_pipeline(cfg, reply="Your notice period is 60 days [user-my-contract]."):
    client = FakeClient(reply)
    return RAGPipeline(cfg, client=client, embedder=Embedder(client, "fake"), log=lambda *_: None), client


def test_clause_headings_detected():
    assert is_clause_heading("10. Termination and Notice Period")
    assert is_clause_heading("Clause 7 - Confidentiality")
    assert is_clause_heading("EMPLOYMENT AGREEMENT")
    assert not is_clause_heading("The Employee will receive a fixed annual salary of INR 18,00,000, paid monthly.")


def test_contract_becomes_clause_sections():
    md = to_markdown("My contract", SAMPLE.read_text(encoding="utf-8"))
    headings = [h for h, _ in split_sections(md)]
    assert "10. Termination and Notice Period" in headings and "4. Compensation" in headings


def test_extract_docx_and_rejects_unknown_type():
    d = docx.Document()
    d.add_heading("Notice Period", level=2)
    d.add_paragraph("Either party may terminate with 60 days notice.")
    buf = io.BytesIO()
    d.save(buf)
    text = extract_text("c.docx", buf.getvalue())
    assert "## Notice Period" in text and "60 days" in text
    with pytest.raises(ValueError):
        extract_text("c.exe", b"x")


def test_save_and_delete_only_user_documents(tmp_path):
    doc_id = save_user_document(tmp_path, "My Contract!", "1. Pay\nINR 10")
    assert doc_id == "user-my-contract"
    assert save_user_document(tmp_path, "My Contract", "x") == "user-my-contract-2"
    assert delete_user_document(tmp_path, doc_id)
    with pytest.raises(ValueError):
        delete_user_document(tmp_path, "hr-leave-policy")


def test_retrieval_restricted_to_uploaded_contract(ucfg):
    save_user_document(ucfg["upload_dir"], "My contract", SAMPLE.read_text(encoding="utf-8"))
    pipe, client = make_pipeline(ucfg)
    r = pipe.answer("What is the notice period after confirmation?", mode="hybrid", doc_ids=["user-my-contract"])
    assert r["sources"] and all(s["doc_id"] == "user-my-contract" for s in r["sources"])
    assert r["citations"] == ["user-my-contract"]
    assert "documents the\nuser added" in client.calls[-1][0]["content"]   # private-document system prompt
    with pytest.raises(ValueError):
        pipe.answer("x?", mode="hybrid", doc_ids=["user-does-not-exist"])


def test_reindex_reuses_existing_embeddings(ucfg):
    pipe, client = make_pipeline(ucfg)
    calls = []
    orig = client.embed
    client.embed = lambda texts, model: (calls.append(len(texts)), orig(texts, model))[1]
    save_user_document(ucfg["upload_dir"], "My contract", SAMPLE.read_text(encoding="utf-8"))
    pipe.reindex(log=lambda *_: None)
    n_contract = sum(1 for c in pipe.index.chunks if c.doc_id == "user-my-contract")
    assert sum(calls) == n_contract        # only the new document was embedded


def test_upload_ask_delete_via_api(ucfg):
    pipe, _ = make_pipeline(ucfg)
    api = TestClient(create_app(pipe))
    r = api.post("/documents/upload", data={"title": "My contract"},
                 files={"file": ("contract.txt", SAMPLE.read_bytes(), "text/plain")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["doc_id"] == "user-my-contract" and body["chunks"] >= 10
    docs = {d["doc_id"]: d for d in api.get("/documents").json()}
    assert docs["user-my-contract"]["uploaded"] and not docs["hr-leave-policy"]["uploaded"]
    a = api.post("/ask", json={"question": "What is my notice period?", "doc_ids": ["user-my-contract"]}).json()
    assert a["doc_filter"] == ["user-my-contract"]
    assert api.post("/documents/upload", data={"title": "empty"}).status_code == 400
    assert api.delete("/documents/hr-leave-policy").status_code == 400
    assert api.delete("/documents/user-my-contract").status_code == 200
    assert "user-my-contract" not in {d["doc_id"] for d in api.get("/documents").json()}
